"""
handlers/settings_handler.py
Displays bot configuration and the contents of the Settings sheet.
"""

from pyrogram import Client, filters
from pyrogram.types import CallbackQuery
from openpyxl import load_workbook

from config import OWNER_ID, EXCEL_FILE, CHANNELS_FILE, APP_NAME
from handlers.access_control import owner_only
from keyboards.main_menu import back_to_menu_keyboard


def register(app: Client) -> None:

    @app.on_callback_query(filters.regex(r"^menu:settings$"))
    @owner_only
    async def settings_menu(client: Client, callback_query: CallbackQuery):
        channels = await client.channel_service.get_all()
        enabled_count = sum(1 for c in channels if c.get("enabled", True))

        settings_rows = []
        try:
            workbook = load_workbook(EXCEL_FILE)
            if "Settings" in workbook.sheetnames:
                sheet = workbook["Settings"]
                for row_idx in range(2, sheet.max_row + 1):
                    key = sheet.cell(row=row_idx, column=1).value
                    value = sheet.cell(row=row_idx, column=2).value
                    if key:
                        settings_rows.append(f"   • {key}: {value}")
        except Exception:
            settings_rows = ["   (could not read Settings sheet)"]

        text = (
            f"⚙ <b>Settings</b>\n\n"
            f"App: {APP_NAME}\n"
            f"Owner ID: <code>{OWNER_ID}</code>\n"
            f"Database file: <code>{EXCEL_FILE.name}</code>\n"
            f"Channels config: <code>{CHANNELS_FILE.name}</code>\n"
            f"Channels: {len(channels)} total, {enabled_count} enabled\n\n"
            f"📋 <b>Settings sheet</b>\n" + ("\n".join(settings_rows) if settings_rows else "   (empty)")
        )
        await callback_query.message.edit_text(text, reply_markup=back_to_menu_keyboard())
        await callback_query.answer()
