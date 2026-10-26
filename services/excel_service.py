"""
services/excel_service.py
Owns database/members.xlsx.

Responsibilities:
    - Auto-create the workbook (with Logs + Settings sheets) if missing.
    - Auto-create a sheet per channel the first time it is needed.
    - Add / remove / unban / search members inside channel sheets.
    - Append operation records into the Logs sheet.
    - Compute statistics used by the Statistics menu.

All public methods are async. openpyxl itself is synchronous, so every
disk-touching call is executed inside an asyncio.Lock to keep writes
serialized and avoid corrupting the workbook when multiple handlers
fire concurrently.
"""

import asyncio
from datetime import datetime
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from config import EXCEL_FILE, SYSTEM_SHEETS, MEMBER_COLUMNS, LOG_COLUMNS
from services.log_service import LogService


class ExcelService:
    def __init__(self, log_service: LogService):
        self._log_service = log_service
        self._lock = asyncio.Lock()
        self._ensure_workbook()

    # ------------------------------------------------------------------
    # Workbook bootstrap
    # ------------------------------------------------------------------
    def _ensure_workbook(self) -> None:
        """Create members.xlsx with Logs + Settings sheets if it does not exist."""
        if EXCEL_FILE.exists():
            return

        workbook = Workbook()
        # openpyxl always creates a default "Sheet" - repurpose it as Logs.
        default_sheet = workbook.active
        default_sheet.title = "Logs"
        default_sheet.append(list(LOG_COLUMNS))

        settings_sheet = workbook.create_sheet("Settings")
        settings_sheet.append(["Key", "Value"])
        settings_sheet.append(["created_at", datetime.now().isoformat()])

        EXCEL_FILE.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(EXCEL_FILE)
        self._log_service.info(f"Created new workbook at {EXCEL_FILE}")

    def _load(self) -> Workbook:
        return load_workbook(EXCEL_FILE)

    def _save(self, workbook: Workbook) -> None:
        workbook.save(EXCEL_FILE)

    def _ensure_channel_sheet(self, workbook: Workbook, channel_name: str) -> Worksheet:
        """Return the sheet for a channel, creating it (with headers) if needed."""
        if channel_name in workbook.sheetnames:
            return workbook[channel_name]

        # Excel sheet titles cannot exceed 31 chars or contain []:*?/\
        safe_name = self._safe_sheet_name(channel_name)
        sheet = workbook.create_sheet(safe_name)
        sheet.append(list(MEMBER_COLUMNS))
        return sheet

    @staticmethod
    def _safe_sheet_name(name: str) -> str:
        invalid = set('[]:*?/\\')
        cleaned = "".join(c for c in name if c not in invalid).strip()
        return cleaned[:31] if cleaned else "Channel"

    def _find_row(self, sheet: Worksheet, telegram_id: Optional[int], username: Optional[str]) -> Optional[int]:
        """Locate the row number (1-indexed, header excluded) matching id or username."""
        username_norm = username.lstrip("@").lower() if username else None
        for row_idx in range(2, sheet.max_row + 1):
            row_id = sheet.cell(row=row_idx, column=1).value
            row_username = sheet.cell(row=row_idx, column=2).value
            if telegram_id is not None and row_id is not None and int(row_id) == int(telegram_id):
                return row_idx
            if username_norm and row_username and str(row_username).lstrip("@").lower() == username_norm:
                return row_idx
        return None

    # ------------------------------------------------------------------
    # Public: channel sheet lifecycle
    # ------------------------------------------------------------------
    async def ensure_channel_sheet(self, channel_name: str) -> None:
        async with self._lock:
            workbook = self._load()
            self._ensure_channel_sheet(workbook, channel_name)
            self._save(workbook)

    async def rename_channel_sheet(self, old_name: str, new_name: str) -> None:
        async with self._lock:
            workbook = self._load()
            if old_name in workbook.sheetnames:
                workbook[old_name].title = self._safe_sheet_name(new_name)
                self._save(workbook)

    async def delete_channel_sheet(self, channel_name: str) -> None:
        async with self._lock:
            workbook = self._load()
            if channel_name in workbook.sheetnames:
                del workbook[channel_name]
                self._save(workbook)

    # ------------------------------------------------------------------
    # Public: member operations
    # ------------------------------------------------------------------
    async def add_member(
        self,
        channel_name: str,
        telegram_id: int,
        username: Optional[str],
        full_name: Optional[str],
        status: str = "Active",
    ) -> bool:
        """Add or update a member row in the given channel's sheet. Returns True if newly added."""
        async with self._lock:
            workbook = self._load()
            sheet = self._ensure_channel_sheet(workbook, channel_name)

            existing_row = self._find_row(sheet, telegram_id, username)
            join_date = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            if existing_row:
                sheet.cell(row=existing_row, column=1, value=telegram_id)
                sheet.cell(row=existing_row, column=2, value=username or "")
                sheet.cell(row=existing_row, column=3, value=full_name or "")
                sheet.cell(row=existing_row, column=5, value=status)
                self._save(workbook)
                return False

            sheet.append([telegram_id, username or "", full_name or "", join_date, status])
            self._save(workbook)
            return True

    async def remove_member(self, channel_name: str, telegram_id: Optional[int] = None,
                             username: Optional[str] = None) -> bool:
        """Delete a member's row from a channel's sheet. Returns True if a row was removed."""
        async with self._lock:
            workbook = self._load()
            if channel_name not in workbook.sheetnames:
                return False
            sheet = workbook[channel_name]
            row = self._find_row(sheet, telegram_id, username)
            if row is None:
                return False
            sheet.delete_rows(row, 1)
            self._save(workbook)
            return True

    async def set_member_status(self, channel_name: str, telegram_id: Optional[int],
                                 username: Optional[str], status: str) -> bool:
        """Update the Status column for a member (used by unban)."""
        async with self._lock:
            workbook = self._load()
            if channel_name not in workbook.sheetnames:
                return False
            sheet = workbook[channel_name]
            row = self._find_row(sheet, telegram_id, username)
            if row is None:
                return False
            sheet.cell(row=row, column=5, value=status)
            self._save(workbook)
            return True

    async def find_member(self, telegram_id: Optional[int] = None,
                           username: Optional[str] = None) -> list[dict]:
        """Search every channel sheet for a member. Returns list of matches with channel name."""
        async with self._lock:
            workbook = self._load()
            matches = []
            for sheet_name in workbook.sheetnames:
                if sheet_name in SYSTEM_SHEETS:
                    continue
                sheet = workbook[sheet_name]
                row = self._find_row(sheet, telegram_id, username)
                if row is not None:
                    matches.append({
                        "channel": sheet_name,
                        "telegram_id": sheet.cell(row=row, column=1).value,
                        "username": sheet.cell(row=row, column=2).value,
                        "full_name": sheet.cell(row=row, column=3).value,
                        "join_date": sheet.cell(row=row, column=4).value,
                        "status": sheet.cell(row=row, column=5).value,
                    })
            return matches

    async def list_channels_with_data(self) -> list[str]:
        async with self._lock:
            workbook = self._load()
            return [s for s in workbook.sheetnames if s not in SYSTEM_SHEETS]

    # ------------------------------------------------------------------
    # Public: logging
    # ------------------------------------------------------------------
    async def write_log(self, action: str, admin: str, user: str, channel: str, status: str) -> None:
        async with self._lock:
            workbook = self._load()
            if "Logs" not in workbook.sheetnames:
                sheet = workbook.create_sheet("Logs")
                sheet.append(list(LOG_COLUMNS))
            else:
                sheet = workbook["Logs"]
            now = datetime.now()
            sheet.append([
                now.strftime("%Y-%m-%d"),
                now.strftime("%H:%M:%S"),
                action,
                admin,
                user,
                channel,
                status,
            ])
            self._save(workbook)
        self._log_service.info(f"[{action}] admin={admin} user={user} channel={channel} status={status}")

    async def get_recent_logs(self, limit: int = 15) -> list[list]:
        async with self._lock:
            workbook = self._load()
            if "Logs" not in workbook.sheetnames:
                return []
            sheet = workbook["Logs"]
            rows = []
            for row_idx in range(2, sheet.max_row + 1):
                rows.append([sheet.cell(row=row_idx, column=c).value for c in range(1, 8)])
            return rows[-limit:][::-1]

    # ------------------------------------------------------------------
    # Public: statistics
    # ------------------------------------------------------------------
    async def get_statistics(self) -> dict:
        async with self._lock:
            workbook = self._load()
            channel_sheets = [s for s in workbook.sheetnames if s not in SYSTEM_SHEETS]

            total_members = 0
            per_channel = {}
            for name in channel_sheets:
                sheet = workbook[name]
                count = max(sheet.max_row - 1, 0)
                per_channel[name] = count
                total_members += count

            today = datetime.now().strftime("%Y-%m-%d")
            today_adds = today_removes = today_unbans = total_removed = 0

            if "Logs" in workbook.sheetnames:
                sheet = workbook["Logs"]
                for row_idx in range(2, sheet.max_row + 1):
                    date_val = sheet.cell(row=row_idx, column=1).value
                    action_val = (sheet.cell(row=row_idx, column=3).value or "").upper()
                    if action_val == "ADD":
                        if date_val == today:
                            today_adds += 1
                    elif action_val == "REMOVE":
                        total_removed += 1
                        if date_val == today:
                            today_removes += 1
                    elif action_val == "UNBAN":
                        if date_val == today:
                            today_unbans += 1

            return {
                "total_channels": len(channel_sheets),
                "total_members": total_members,
                "total_removed": total_removed,
                "members_per_channel": per_channel,
                "today_adds": today_adds,
                "today_removes": today_removes,
                "today_unbans": today_unbans,
            }
